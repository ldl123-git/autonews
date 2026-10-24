"""表格解析：支持 xlsx / xls / csv，流式读取，边读边入库。"""
import csv
import datetime as dt
import re
import sqlite3
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional

from . import style
from .config import IMPORT_COMMIT_SIZE
from .db import get_conn, write_conn

FIELD_ALIASES: Dict[str, List[str]] = {
    "title": ["标题", "新闻标题", "题目", "标题名", "title"],
    "source": ["来源", "平台", "站点", "渠道", "来源平台", "source"],
    "orig_category": ["分类", "原始分类", "频道", "栏目", "类别", "category"],
    "author": ["作者", "账号", "作者名", "author"],
    "content_type": ["类型", "内容类型", "形式", "type", "content_type"],
    "publish_time": ["发布时间", "时间", "发布日期", "publish_time", "date", "publish"],
    "read_count": ["阅读", "阅读量", "阅读数", "浏览量", "read", "read_count", "views"],
    "like_count": ["点赞", "点赞数", "点赞量", "like", "likes", "like_count"],
    "comment_count": ["评论", "评论数", "评论量", "comment", "comments"],
    "share_count": ["分享", "分享数", "转发", "转发数", "share"],
    "url": ["链接", "原文链接", "地址", "原文地址", "url", "link"],
}

NUM_FIELDS = ("read_count", "like_count", "comment_count", "share_count")


def _norm_key(name: Any) -> str:
    s = str(name or "").strip().lower()
    s = re.sub(r"[\s_\-（）()\[\]【】：:]", "", s)
    return s


def build_field_map(header: List[Any]) -> Dict[int, str]:
    """把表头列索引映射到内部字段名。"""
    mapping: Dict[int, str] = {}
    for idx, raw in enumerate(header):
        key = _norm_key(raw)
        if not key:
            continue
        for field, aliases in FIELD_ALIASES.items():
            if key in [_norm_key(a) for a in aliases]:
                mapping[idx] = field
                break
        else:
            for field, aliases in FIELD_ALIASES.items():
                if any(_norm_key(a) in key or key in _norm_key(a) for a in aliases if len(_norm_key(a)) > 1):
                    mapping[idx] = field
                    break
    return mapping


def parse_count(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).strip().replace(",", "").replace("，", "")
    if not s:
        return 0
    m = re.match(r"^(\d+(?:\.\d+)?)\s*([万wWkK千]?)\+?$", s)
    if m:
        num = float(m.group(1))
        unit = m.group(2)
        factor = {"万": 10000, "w": 10000, "W": 10000, "k": 1000, "K": 1000, "千": 1000}.get(unit, 1)
        return int(num * factor)
    digits = re.sub(r"[^\d]", "", s)
    return int(digits) if digits else 0


def _from_excel_serial(value: float) -> Optional[dt.datetime]:
    try:
        if value > 100000:  # 毫秒时间戳
            return dt.datetime.fromtimestamp(value / 1000)
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))
    except Exception:
        return None


TIME_FORMATS = [
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y/%m/%d %H:%M:%S",
    "%Y/%m/%d %H:%M",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y年%m月%d日 %H:%M",
    "%Y年%m月%d日",
    "%Y.%m.%d",
]


def parse_time(value: Any) -> str:
    """统一输出 YYYY-MM-DD HH:MM:SS，解析失败返回空串。"""
    if value is None or value == "":
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d 00:00:00")
    if isinstance(value, (int, float)):
        d = _from_excel_serial(float(value))
        return d.strftime("%Y-%m-%d %H:%M:%S") if d else ""
    s = str(value).strip()
    if not s:
        return ""
    for fmt in TIME_FORMATS:
        try:
            return dt.datetime.strptime(s, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    m = re.search(r"(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})(?:[日\s]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        hh = int(m.group(4) or 0)
        mi = int(m.group(5) or 0)
        ss = int(m.group(6) or 0)
        try:
            return dt.datetime(y, mo, d, hh, mi, ss).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return ""
    return ""


def normalize_content_type(value: Any) -> str:
    s = str(value or "").strip()
    if not s:
        return ""
    if "视频" in s or "video" in s.lower():
        return "视频"
    if "图" in s or "image" in s.lower() or "pic" in s.lower():
        return "图片"
    if "文" in s or "article" in s.lower():
        return "文章"
    return s


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def iter_rows(path: Path) -> Iterator[Dict[str, Any]]:
    """按扩展名选择解析器，产出归一化的行字典（含表头行之前会自动跳过）。"""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        yield from _iter_xlsx(path)
    elif suffix == ".xls":
        yield from _iter_xls(path)
    else:
        yield from _iter_csv(path)


def _iter_xlsx(path: Path) -> Iterator[Dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        field_map: Dict[int, str] = {}
        for row in ws.iter_rows(values_only=True):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            if not field_map:
                field_map = build_field_map(list(row))
                if not field_map:
                    raise ValueError("未能识别表头，请确认表格首行包含：标题 / 来源 / 分类 / 发布时间 等列")
                continue
            yield _build_row(field_map, row)
    finally:
        wb.close()


def _iter_xls(path: Path) -> Iterator[Dict[str, Any]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(0)
    field_map: Dict[int, str] = {}
    for idx in range(sheet.nrows):
        row = [sheet.cell_value(idx, c) for c in range(sheet.ncols)]
        if all(str(c or "").strip() == "" for c in row):
            continue
        if not field_map:
            field_map = build_field_map(row)
            if not field_map:
                raise ValueError("未能识别表头，请确认表格首行包含：标题 / 来源 / 分类 / 发布时间 等列")
            continue
        yield _build_row(field_map, row)


def _iter_csv(path: Path) -> Iterator[Dict[str, Any]]:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别")
    reader = csv.reader(text.splitlines())
    field_map: Dict[int, str] = {}
    for row in reader:
        if not row or all(str(c or "").strip() == "" for c in row):
            continue
        if not field_map:
            field_map = build_field_map(row)
            if not field_map:
                raise ValueError("未能识别表头，请确认表格首行包含：标题 / 来源 / 分类 / 发布时间 等列")
            continue
        yield _build_row(field_map, row)


def _build_row(field_map: Dict[int, str], row: List[Any]) -> Dict[str, Any]:
    data: Dict[str, Any] = {
        "title": "", "source": "", "author": "", "orig_category": "", "content_type": "",
        "publish_time": "", "url": "", "read_count": 0, "like_count": 0,
        "comment_count": 0, "share_count": 0,
    }
    for idx, field in field_map.items():
        if idx >= len(row):
            continue
        value = row[idx]
        if field in NUM_FIELDS:
            data[field] = parse_count(value)
        elif field == "publish_time":
            data[field] = parse_time(value)
        elif field == "content_type":
            data[field] = normalize_content_type(value)
        else:
            data[field] = _text(value)
    return data


def create_batch(filename: str, size: int) -> int:
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with write_conn() as conn:
        cur = conn.execute(
            "INSERT INTO batches (filename, size, status, created_at) VALUES (?,?, 'parsing', ?)",
            (filename, size, now),
        )
        return int(cur.lastrowid)


def run_import(batch_id: int, path: Path) -> None:
    """后台执行：流式解析并写入 articles，实时更新批次进度。"""
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    inserted = duplicated = total = 0
    buffer: List[tuple] = []
    try:
        with write_conn() as conn:
            for row in iter_rows(path):
                total += 1
                if not row["title"] and not row["url"]:
                    continue
                publish_time = row["publish_time"]
                publish_date = publish_time[:10]
                buffer.append((
                    batch_id, row["title"], row["source"],
                    style.normalize_source(row["source"]),   # 来源标记：导入时就归一
                    row["author"], row["orig_category"],
                    row["orig_category"], row["content_type"], row["read_count"], row["like_count"],
                    row["comment_count"], row["share_count"], row["url"], publish_time, publish_date,
                    now, now,
                ))
                if len(buffer) >= IMPORT_COMMIT_SIZE:
                    ok, skip = _flush(conn, buffer)
                    inserted += ok
                    duplicated += skip
                    buffer.clear()
                    _update_progress(conn, batch_id, total, inserted, duplicated)
            if buffer:
                ok, skip = _flush(conn, buffer)
                inserted += ok
                duplicated += skip
                buffer.clear()
            _update_progress(conn, batch_id, total, inserted, duplicated)
            conn.execute(
                "UPDATE batches SET status='done', message=? WHERE id=?",
                (f"导入完成：读取 {total} 行，新增 {inserted} 条，跳过重复 {duplicated} 条", batch_id),
            )
    except Exception as exc:  # noqa: BLE001
        with write_conn() as conn:
            conn.execute(
                "UPDATE batches SET status='error', message=? WHERE id=?",
                (f"导入失败：{exc}", batch_id),
            )


def _flush(conn: sqlite3.Connection, buffer: List[tuple]) -> tuple:
    """返回 (成功写入条数, 被去重跳过条数)。"""
    sql = """
        INSERT OR IGNORE INTO articles
        (batch_id, title, source, source_norm, author, orig_category, category, content_type,
         read_count, like_count, comment_count, share_count, url, publish_time,
         publish_date, created_at, updated_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """
    cur = conn.executemany(sql, buffer)
    ok = max(cur.rowcount or 0, 0)
    return ok, len(buffer) - ok


def _update_progress(conn: sqlite3.Connection, batch_id: int, total: int, inserted: int, duplicated: int) -> None:
    conn.execute(
        "UPDATE batches SET total=?, inserted=?, duplicated=? WHERE id=?",
        (total, inserted, duplicated, batch_id),
    )


def list_batches() -> List[Dict[str, Any]]:
    with get_conn() as conn:
        rows = conn.execute("SELECT * FROM batches ORDER BY id DESC LIMIT 50").fetchall()
        return [dict(r) for r in rows]


def get_batch(batch_id: int) -> Optional[Dict[str, Any]]:
    with get_conn() as conn:
        row = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
        return dict(row) if row else None
